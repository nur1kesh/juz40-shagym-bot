"""Бизнес-логика сценарийлері: жалған Telegram + басқарылатын сағат. Нақты желі/OpenAI керек емес."""
import asyncio
import io
import json
import os
import tempfile
from datetime import datetime, timedelta
from types import SimpleNamespace as NS
from unittest.mock import AsyncMock

import pytest

os.environ.update(BOT_TOKEN="1:x", OPENAI_API_KEY="sk-x",
                  DB_PATH=os.path.join(tempfile.mkdtemp(), "t.db"))

from bot import config, db, exporter, service, templates as T  # noqa: E402
import bot.main as bm  # noqa: E402

REAL_WINDOW = config.in_work_window
from bot import ai as _ai_mod  # noqa: E402
REAL_CLASSIFY = _ai_mod.classify
SMART, JUNIOR, EXT = (-100, 5), (-100, 6), (-200, 9)


def run(coro):
    return asyncio.run(coro)


class Clock:
    def __init__(self, dt):
        self.dt = dt

    def now(self):
        return self.dt

    def set(self, y, m, d, h=12, mi=0):
        self.dt = datetime(y, m, d, h, mi, tzinfo=config.TZ)

    def add(self, **kw):
        self.dt += timedelta(**kw)


class FakeBot:
    def __init__(self):
        self.sent, self.copied, self.reactions, self.edited = [], [], [], []
        self.fail_chat = None
        self.deleted = []
        self._id = 5000

    async def send_message(self, chat_id, text, **kw):
        if self.fail_chat == chat_id:
            raise RuntimeError("boom")
        self._id += 1
        self.sent.append(NS(chat=chat_id, thread=kw.get("message_thread_id"), text=text,
                            reply_to=kw.get("reply_to_message_id"), kb=kw.get("reply_markup"), id=self._id))
        return NS(message_id=self._id)

    async def copy_messages(self, chat_id, from_chat_id, ids, **kw):
        self.copied.append((chat_id, list(ids)))
        out = []
        for _ in ids:
            self._id += 1
            out.append(NS(message_id=self._id))
        return out

    async def set_message_reaction(self, *a, **k):
        self.reactions.append(a)

    async def delete_message(self, chat_id, msg_id):
        self.deleted.append((chat_id, msg_id))

    async def download(self, file, destination=None):
        return io.BytesIO(b"\xff\xd8fake-jpeg")

    async def edit_message_text(self, text, chat_id=None, message_id=None, **kw):
        self.edited.append((chat_id, message_id, text))

    def texts(self):
        return [m.text for m in self.sent]


@pytest.fixture(autouse=True)
def env(monkeypatch):
    for t in ("bindings", "users", "complaints", "replies", "bot_msgs", "notices", "pending_res",
              "media_msgs", "media_groups", "ephemeral", "digest_msgs", "deletions", "errors", "decisions"):
        db.x(f"DELETE FROM {t}")
    db.bind("smart_hate", *SMART)
    db.bind("junior_hate", *JUNIOR)
    db.bind("external_hate", *EXT)
    clock = Clock(datetime(2026, 10, 14, 12, 0, tzinfo=config.TZ))  # сәрсенбі
    monkeypatch.setattr(config, "now", clock.now)
    monkeypatch.setattr(service, "SEND_GAP", 0)
    monkeypatch.setattr(service, "now", clock.now)
    monkeypatch.setattr(service, "in_work_window", lambda dt=None: REAL_WINDOW(dt or clock.now()))
    monkeypatch.setattr(service.ai, "check_solution", AsyncMock(return_value={"ok": True, "reason": ""}))
    monkeypatch.setattr(service.ai, "extract_complaint", AsyncMock(return_value=None))
    monkeypatch.setattr(service.ai, "read_screenshot", AsyncMock(return_value=None))
    monkeypatch.setattr(service.ai, "classify_reply", AsyncMock(return_value={"kind": "other"}))
    monkeypatch.setattr(service.ai, "classify", AsyncMock(
        return_value={"role": "Куратор", "title": "Кеш жауап беру", "employee_name": "Есболқызы Арай"}))
    return NS(clock=clock, bot=FakeBot())


_mid = [0]
_users = {}


def U(uid, name, username=None):
    u = NS(id=uid, full_name=name, username=username, is_bot=False)
    db.remember_user(u)
    return u


TECH = lambda: U(1, "Бекарыс Алишер | қолдау")  # noqa: E731  (тех.менеджерлер тізімінде бар)
COORD = lambda: U(2, "Бақыт Аружан | ГЕО", "aruzhanbakyt")  # noqa: E731
EXTSELL = lambda: U(3, "Айтикова Аружан | сату")  # noqa: E731


def M(user, where, text=None, caption=None, reply=None, group=None, ctype="text"):
    _mid[0] += 1
    chat, thread = where
    rt = NS(message_id=reply, forum_topic_created=None) if reply else None
    return NS(text=text, caption=caption, content_type=ctype, media_group_id=group, chat=NS(id=chat),
              message_id=_mid[0], message_thread_id=thread, is_topic_message=True, from_user=user,
              reply_to_message=rt)


def complaint_text(student="Айдана Нұрлан", subject="SMART ГЕО-31", title="Кеш жауап беру", desc="Куратор жауап бермейді"):
    return (f"Оқушы аты-жөні: {student}\nТелефон нөмірі: 87001112233\nПәні және потогы: {subject}\n"
            f"Қызметкердің аты-жөні: Есболқызы Арай\nҚызметі/Мәселе иесі: Куратор\nМәселе атауы/типі: {title}\n"
            f"Мәселенің толық сипаттамасы: {desc}")


def resolution_text(student="Айдана Нұрлан", just="Негізделген", sat="Иә", stay="Иә"):
    return (f"Оқушы аты-жөні: {student}\nМәселенің шешімі: Оқушымен байланысқа шықтық, себебі анықталды, "
            f"кураторға ескерту жасалды.\nМәселе шындыққа негізделген бе? {just}\nОқушы риза ма? {sat}\n"
            f"Оқушы курста қалды ма? {stay}")


def C():
    return db.q("SELECT * FROM complaints ORDER BY id")


def post(e, user, where, text, **kw):
    msg = M(user, where, text=text, **kw)
    run(bm.handle(msg, e.bot, False))
    return msg


# ===================================================================== жұмыс терезесі
@pytest.mark.parametrize("args,expected", [
    ((2026, 10, 14, 10, 59), False), ((2026, 10, 14, 11, 0), True), ((2026, 10, 14, 20, 59), True),
    ((2026, 10, 14, 21, 0), False), ((2026, 10, 17, 15, 0), True),   # сенбі - жұмыс күні
    ((2026, 10, 18, 15, 0), False),                                  # жексенбі - демалыс
])
def test_work_window(args, expected):
    y, m, d, h, mi = args
    assert REAL_WINDOW(datetime(y, m, d, h, mi, tzinfo=config.TZ)) is expected


# ===================================================================== негізгі ағын
def test_direct_flow_end_to_end(env):
    tech = TECH()
    m = post(env, tech, SMART, complaint_text())
    c = C()
    assert len(c) == 1 and c[0]["status"] == "open" and c[0]["program"] == "smart"
    assert c[0]["author_role"] == "tech"                       # аты тізімнен танылды
    notify = env.bot.sent[0]
    assert notify.reply_to == m.message_id and (notify.chat, notify.thread) == SMART
    for tag in ("@zhhaniya", "@aruzhanbakyt", "@SdlyAmina", "@Meyrambekovich"):
        assert tag in notify.text
    assert env.bot.reactions                                   # 👀

    # бір сағаттан кейін еске салу
    env.clock.add(minutes=59); run(service.flush(env.bot)); assert len(env.bot.sent) == 1
    env.clock.add(minutes=2); run(service.flush(env.bot)); assert len(env.bot.sent) == 2
    assert "Еске салу" in env.bot.sent[1].text and "@zhhaniya" in env.bot.sent[1].text

    # координатор reply ішінде шешім береді
    post(env, COORD(), SMART, resolution_text(), reply=m.message_id)
    c = C()[0]
    assert c["status"] == "resolved"
    assert json.loads(c["final"])["satisfied"] == "Иә"
    assert "шешілді" in env.bot.sent[-1].text
    n = len(env.bot.sent)
    env.clock.add(hours=3); run(service.flush(env.bot)); assert len(env.bot.sent) == n   # еске салу тоқтады

    xl = __import__("openpyxl").load_workbook(io.BytesIO(exporter.build()))["ТЕХ МЕНЕДЖЕР"]
    row = [x.value for x in xl[2]]
    assert row[1] == "Айдана Нұрлан" and row[4] == "Куратор" and row[5] == "Кеш жауап беру"
    assert row[8] == "Негізделген" and row[10] == "Бекарыс Алишер" and row[11] == "Иә" and row[12] is True


def test_reply_to_bot_message_also_resolves(env):
    post(env, TECH(), SMART, complaint_text())
    notify = env.bot.sent[0]
    post(env, COORD(), SMART, resolution_text(), reply=notify.id)
    assert C()[0]["status"] == "resolved"


# ===================================================================== жұмыс уақыты / SLA
def test_night_complaint_is_handled_immediately_but_reminders_wait_for_the_window(env):
    env.clock.set(2026, 10, 17, 22, 30)                       # сенбі түн
    post(env, TECH(), SMART, complaint_text())
    assert C()[0]["status"] == "open" and len(env.bot.sent) == 1          # бірден қабылданып, тегталды
    assert "21.10" not in env.bot.sent[0].text and "18.10 22:30" in env.bot.sent[0].text   # дедлайн = +24 сағ
    env.clock.set(2026, 10, 18, 14, 0); run(service.flush(env.bot))       # жексенбі: еске салу жоқ
    assert len(env.bot.sent) == 1
    env.clock.set(2026, 10, 19, 11, 0); run(service.flush(env.bot))       # дүйсенбі 11:00: бірінші еске салу
    assert len(env.bot.sent) == 2 and "Дедлайн өтті" in env.bot.sent[1].text      # 24 сағат күнтізбелік: сенбі 22:30 + 24 сағ өтіп кеткен


def test_reminders_hourly_only_in_window_and_overdue_marker(env):
    post(env, TECH(), SMART, complaint_text())                # сәрсенбі 12:00
    for _ in range(9):                                        # 13:00 ... 21:00 дейін
        env.clock.add(hours=1)
        run(service.flush(env.bot))
    # 13..20 = 8 еске салу, 21:00-де терезе жабық
    assert len([t for t in env.bot.texts() if "Еске салу" in t]) == 8
    before = len(env.bot.sent)
    env.clock.set(2026, 10, 15, 3, 0); run(service.flush(env.bot)); assert len(env.bot.sent) == before
    env.clock.set(2026, 10, 15, 12, 30); run(service.flush(env.bot))
    assert "Дедлайн өтті" in env.bot.sent[-1].text


def test_sunday_complaint_accepted_reminders_only_monday_and_not_a_flood(env):
    env.clock.set(2026, 10, 18, 12, 0)                        # жексенбі
    post(env, TECH(), SMART, complaint_text())
    assert len(env.bot.sent) == 1                             # шағым жексенбіде де қабылданады
    for _ in range(5):
        env.clock.add(hours=1); run(service.flush(env.bot))
    assert len(env.bot.sent) == 1                             # жексенбіде еске салу жоқ
    env.clock.set(2026, 10, 19, 11, 0); run(service.flush(env.bot)); run(service.flush(env.bot))
    assert len(env.bot.sent) == 2                             # дүйсенбіде бір ғана еске салу


# ===================================================================== шаблон тексеру
def test_invalid_complaint_gets_error_then_edit_registers_it(env):
    bad = complaint_text(subject="SMART пәні жоқ")            # пән танылмайды
    msg = post(env, TECH(), SMART, bad)
    assert C() == [] and "танылмады" in env.bot.sent[0].text
    msg.text = complaint_text()
    run(bm.handle(msg, env.bot, True))
    assert len(C()) == 1 and C()[0]["status"] == "open"


def test_errors_and_resolution_work_outside_the_window_too(env):
    env.clock.set(2026, 10, 18, 12, 0)                        # жексенбі
    msg = post(env, TECH(), SMART, complaint_text(subject="SMART пәні жоқ"))
    assert "Түзетіңіз" in env.bot.sent[0].text                # қате бірден айтылады
    msg.text = complaint_text()
    run(bm.handle(msg, env.bot, True))
    assert len(C()) == 1 and C()[0]["status"] == "open"
    post(env, COORD(), SMART, "в процессе", reply=env.bot.sent[-1].id)
    assert C()[0]["progress_at"] and env.bot.reactions[-1][2][0].emoji == "✍"    # реакция да жексенбіде
    post(env, COORD(), SMART, resolution_text(), reply=env.bot.sent[-1].id)
    assert C()[0]["status"] == "resolved" and "шешілді" in env.bot.sent[-1].text


def test_screenshot_without_text_prompts_template(env):
    post(env, TECH(), SMART, None, ctype="photo")
    assert "Мәтінсіз" in env.bot.sent[0].text and C() == []


def test_chat_noise_is_ignored(env):
    post(env, TECH(), SMART, "рахмет")
    post(env, TECH(), SMART, "Қайырлы күн, бүгін сабақ қай уақытта басталады екен деп сұрап жатыр, кім біледі?" * 1)
    assert C() == []
    assert len(env.bot.sent) == 1 and "Шаблонсыз" in env.bot.sent[0].text   # тек ұзынына ескерту


def test_roster_member_message_is_not_a_complaint(env):
    post(env, COORD(), SMART, complaint_text())
    assert C() == []


def test_program_mismatch_in_topic_is_rejected(env):
    post(env, TECH(), JUNIOR, complaint_text(subject="SMART ГЕО-31"))
    assert C() == [] and "SMART" in env.bot.sent[0].text and "JUNIOR" in env.bot.sent[0].text


def test_duplicate_open_complaint_not_created_twice(env):
    t = TECH()
    post(env, t, SMART, complaint_text())
    post(env, t, SMART, complaint_text())
    assert len(C()) == 1


# ===================================================================== кім шеше алады
def test_author_cannot_close_own_complaint(env):
    tech = TECH()
    m = post(env, tech, SMART, complaint_text())
    post(env, tech, SMART, resolution_text(), reply=m.message_id)
    assert C()[0]["status"] == "open"


def test_other_tech_manager_cannot_close(env):
    m = post(env, TECH(), SMART, complaint_text())
    other = U(11, "Куанов Али | қолдау")
    post(env, other, SMART, resolution_text(), reply=m.message_id)
    assert C()[0]["status"] == "open"


def test_resolution_validation_errors_keep_complaint_open_and_edit_fixes(env):
    m = post(env, TECH(), SMART, complaint_text())
    r = post(env, COORD(), SMART, resolution_text(sat="мүмкін"), reply=m.message_id)
    assert C()[0]["status"] == "open" and "Риза ма" in env.bot.sent[-1].text
    r.text = resolution_text()
    run(bm.handle(r, env.bot, True))
    assert C()[0]["status"] == "resolved"


def test_vague_solution_rejected_by_ai_check(env, monkeypatch):
    monkeypatch.setattr(service.ai, "check_solution",
                        AsyncMock(return_value={"ok": False, "reason": "нақты емес"}))
    m = post(env, TECH(), SMART, complaint_text())
    post(env, COORD(), SMART, resolution_text(), reply=m.message_id)
    assert C()[0]["status"] == "open" and "нақты емес" in env.bot.sent[-1].text


# ===================================================================== reply-сыз шешім
def test_loose_resolution_single_match_closes(env):
    post(env, TECH(), SMART, complaint_text())
    post(env, COORD(), SMART, resolution_text())
    assert C()[0]["status"] == "resolved"


def test_loose_resolution_multiple_matches_asks_and_pick_closes(env):
    post(env, TECH(), SMART, complaint_text(student="Айдана Нұрлан"))
    post(env, TECH(), SMART, complaint_text(student="Айдана Сәрсен", subject="SMART БИО-11", desc="басқа мәселе"))
    post(env, COORD(), SMART, resolution_text(student="Айдана Нұрлан").replace("Айдана Нұрлан", "Айдана"))
    q = env.bot.sent[-1]
    assert "Қай шағымға" in q.text
    cb = q.kb.inline_keyboard[0][0].callback_data
    pid, cid = cb.split(":")[1:]
    msg = run(service.finish_pending(env.bot, int(pid), int(cid)))
    assert "шешілді" in msg and [c["status"] for c in C()].count("resolved") == 1


def test_loose_resolution_unknown_student(env):
    post(env, TECH(), SMART, complaint_text())
    post(env, COORD(), SMART, resolution_text(student="Бейтаныс Адам"))
    assert C()[0]["status"] == "open" and "жоқ" in env.bot.sent[-1].text


# ===================================================================== сыртқы сату
def test_external_flow_forward_and_echo(env):
    ext = EXTSELL()
    src = post(env, ext, EXT, complaint_text())
    c = C()[0]
    assert c["source"] == "external" and c["status"] == "open" and c["program"] == "smart"
    assert env.bot.copied == [(SMART[0], [src.message_id])]
    assert (env.bot.sent[0].chat, env.bot.sent[0].thread) == SMART
    copied_id = c["target_msg"]
    post(env, COORD(), SMART, resolution_text(), reply=copied_id)
    assert C()[0]["status"] == "resolved"
    dests = [(m.chat, m.thread) for m in env.bot.sent if "шешілді" in m.text]
    assert dests == [SMART, EXT]                               # екі жаққа да жауап
    assert env.bot.sent[-1].reply_to == src.message_id
    xl = __import__("openpyxl").load_workbook(io.BytesIO(exporter.build()))
    assert "СЫРТҚЫ САТУ" in xl.sheetnames and xl["СЫРТҚЫ САТУ"][2][11].value == "Айтикова Аружан"


def test_external_without_program_asks_buttons_and_pick_routes(env):
    ext = EXTSELL()
    post(env, ext, EXT, complaint_text(subject="ГЕО-31"))
    c = C()[0]
    assert c["status"] == "awaiting_program"
    assert "SMART па" in env.bot.sent[0].text
    tgt = db.binding("smart_hate")
    parsed = json.loads(c["parsed"])
    coords, sups, notes = service.compute_responsible("smart", parsed)
    db.update("complaints", c["id"], program="smart", status="new", last_reminder_at=None, target_chat=tgt["chat_id"],
              target_thread=tgt["thread_id"], responsible=service.resp_json(coords, sups), notes="")
    run(service.flush(env.bot))
    assert C()[0]["status"] == "open" and env.bot.copied


def test_junior_routing(env):
    post(env, TECH(), JUNIOR, complaint_text(subject="JUNIOR МАТ"))
    t = env.bot.sent[0].text
    assert "@arystanbekova7" in t and "@maqzhan1" in t and "@bayashkarr" not in t
    post(env, TECH(), JUNIOR, complaint_text(student="Басқа Оқушы", subject="JUNIOR БИО", desc="x"))
    t = env.bot.sent[1].text
    assert "@bayashkarr" in t and "@nurlnvash" in t and "@maqzhan1" in t


# ===================================================================== сенімділік
def test_concurrent_flush_does_not_duplicate_messages(env):
    tech = TECH()
    msg = M(tech, SMART, text=complaint_text())
    db.x("INSERT INTO complaints(status,source,src_chat,src_thread,src_msg,author_id,author_name,author_role,program,"
         "target_chat,target_thread,target_msg,text,parsed,responsible,created_at) VALUES('new','direct',?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
         SMART[0], SMART[1], msg.message_id, 1, "x", "tech", "smart", SMART[0], SMART[1], msg.message_id, "t",
         json.dumps(T.validate_complaint(complaint_text())[0]), json.dumps({}), env.clock.now().isoformat())

    async def both():
        await asyncio.gather(service.flush(env.bot), service.flush(env.bot), service.flush(env.bot))
    run(both())
    assert len([t for t in env.bot.texts() if "🔔" in t]) == 1


def test_echo_failure_does_not_repeat_target_notice(env):
    post(env, EXTSELL(), EXT, complaint_text())
    c = C()[0]
    env.bot.fail_chat = EXT[0]                                 # сыртқы топқа жазу сәтсіз
    post(env, COORD(), SMART, resolution_text(), reply=c["target_msg"])
    for _ in range(3):
        env.clock.add(seconds=20); run(service.flush(env.bot))
    assert len([m for m in env.bot.sent if m.chat == SMART[0] and "шешілді" in m.text]) == 1


def test_dispatch_failure_after_copy_does_not_copy_again(env):
    env.bot.fail_chat = SMART[0]                               # notify жіберу сәтсіз, copy сәтті
    post(env, EXTSELL(), EXT, complaint_text())
    env.bot.fail_chat = None
    for _ in range(3):
        env.clock.add(minutes=6); run(service.flush(env.bot))
    assert len(env.bot.copied) == 1


def test_excel_formula_injection_is_neutralised(env):
    m = post(env, TECH(), SMART, complaint_text(desc="=HYPERLINK(\"http://evil\",\"x\")"))
    post(env, COORD(), SMART, resolution_text(), reply=m.message_id)
    ws = __import__("openpyxl").load_workbook(io.BytesIO(exporter.build()))["ТЕХ МЕНЕДЖЕР"]
    cell = ws[2][6]
    assert cell.data_type != "f"


# ===================================================================== пән талдау
@pytest.mark.parametrize("label,program,pairs", [
    ("SMART ГЕО-31", "smart", [("ГЕО", "31")]),
    ("SMART ФИЗ, МАТ-01", "smart", [("ФИЗ", "01"), ("МАТ", "01")]),
    ("SMART ДЖТ-11, АНГЛ-11, МС-71, Т-71", "smart", [("ДЖТ", "11"), ("АНГЛ", "11"), ("МС", "71"), ("ТАРИХ", "71")]),
    ("JUNIOR БИО", "junior", [("БИО", None)]),
    ("SMART Орыс тілі-21", "smart", [("РУСЛИТ", "21")]),
    ("SMART Ағылшын тілі-01", "smart", [("АНГЛ", "01")]),
    ("SMART Қазақстан тарихы-31", "smart", [("ТАРИХ", "31")]),
    ("СМАРТ МАТ-11", "smart", [("МАТ", "11")]),
    ("Джуниор БИО", "junior", [("БИО", None)]),
])
def test_subject_parsing(label, program, pairs):
    p, got = T.parse_subject(label)
    assert p == program and got == pairs


# ===================================================================== чат таза болуы (қысқа хабарлар)
def test_bot_messages_are_short(env):
    post(env, TECH(), SMART, "Қайырлы күн, бүгін сабақ қай уақытта басталады екен деп сұрап жатыр, кім біледі?")
    post(env, TECH(), SMART, None, ctype="photo")
    post(env, TECH(), SMART, complaint_text(subject="SMART пәні жоқ"))        # қате
    post(env, TECH(), SMART, complaint_text(student="Басқа Бала", desc="x"))   # дұрыс
    env.clock.add(hours=1); run(service.flush(env.bot))                  # еске салу
    for m in env.bot.sent:
        assert len(m.text) < 260, (len(m.text), m.text)


def test_error_notice_is_deleted_when_fixed(env):
    msg = post(env, TECH(), SMART, complaint_text(subject="SMART пәні жоқ"))
    err = env.bot.sent[0]
    msg.text = complaint_text()
    run(bm.handle(msg, env.bot, True))
    assert (err.chat, err.id) in env.bot.deleted


def test_notice_auto_deletes_after_ttl(env):
    post(env, TECH(), SMART, None, ctype="photo")
    n = env.bot.sent[0]
    env.clock.add(minutes=10); run(service.flush(env.bot)); assert (n.chat, n.id) not in env.bot.deleted
    env.clock.add(minutes=6); run(service.flush(env.bot)); assert (n.chat, n.id) in env.bot.deleted


def test_reminders_are_digest_when_many_open(env):
    for i in range(3):
        post(env, TECH(), SMART, complaint_text(student=f"Оқушы Нөмір{i}", desc=f"мәселе {i}"))
    sent_before = len(env.bot.sent)
    env.clock.add(hours=1); run(service.flush(env.bot))
    new = env.bot.sent[sent_before:]
    assert len(new) == 1, [m.text for m in new]                         # 3 емес, 1 хабарлама
    assert "Шешілмеген шағымдар (3)" in new[0].text and new[0].text.count("t.me/c/") == 3


def test_resolved_notice_short_in_topic_full_echo_for_external(env):
    ext = post(env, EXTSELL(), EXT, complaint_text())
    post(env, COORD(), SMART, resolution_text(), reply=C()[0]["target_msg"])
    topic_msg = [m for m in env.bot.sent if m.chat == SMART[0] and "шешілді" in m.text][0]
    echo_msg = [m for m in env.bot.sent if m.chat == EXT[0] and "шешілді" in m.text][0]
    assert len(topic_msg.text) < 80 and "Оқушымен байланысқа" not in topic_msg.text
    assert "Оқушымен байланысқа" in echo_msg.text


# ===================================================================== автоматты іріктеу (OpenAI)
def minimal_text(student="Айдана Нұрлан", subject="SMART ГЕО-31", desc="Куратор Арай 3 күн жауап бермейді"):
    return (f"Оқушы аты-жөні: {student}\nПәні және потогы: {subject}\n"
            f"Мәселенің толық сипаттамасы: {desc}")


def test_template_has_no_prefill_and_no_option_lists():
    assert "SMART ГЕО-31" not in T.COMPLAINT_TEMPLATE and "SMART ГЕО-31" not in T.COMPLAINT_HELP
    assert "Қызметі" not in T.COMPLAINT_HELP and "Мәселе типі" not in T.COMPLAINT_HELP
    assert "blockquote" not in T.COMPLAINT_HELP


def test_minimal_complaint_gets_role_title_employee_from_ai(env, monkeypatch):
    monkeypatch.setattr(service.ai, "classify", AsyncMock(
        return_value={"role": "Мұғалім", "title": "Мұғалім мәселесі", "employee_name": "Асан Бекұлы"}))
    post(env, TECH(), SMART, minimal_text())
    p = json.loads(C()[0]["parsed"])
    assert (p["employee_role"], p["issue_title"], p["employee_name"]) == ("Мұғалім", "Мұғалім мәселесі", "Асан Бекұлы")
    assert "Мұғалім мәселесі" in env.bot.sent[0].text


def test_manual_values_are_kept_and_ai_not_called(env, monkeypatch):
    spy = AsyncMock(return_value={"role": "Курс", "title": "Курс системасы", "employee_name": "X"})
    monkeypatch.setattr(service.ai, "classify", spy)
    post(env, TECH(), SMART, complaint_text())          # қызметі/типі/қызметкер қолмен жазылған
    p = json.loads(C()[0]["parsed"])
    assert (p["employee_role"], p["issue_title"], p["employee_name"]) == ("Куратор", "Кеш жауап беру", "Есболқызы Арай")
    spy.assert_not_awaited()


def test_invalid_manual_role_is_not_an_error_ai_decides(env):
    text = complaint_text().replace("Қызметі/Мәселе иесі: Куратор", "Қызметі/Мәселе иесі: ойдан шығарылған")
    post(env, TECH(), SMART, text)
    assert len(C()) == 1 and json.loads(C()[0]["parsed"])["employee_role"] == "Куратор"   # AI (mock) таңдады


def test_classify_validates_ai_output_and_falls_back_to_dash(monkeypatch):
    from bot import ai

    def resp(content):
        return NS(choices=[NS(message=NS(content=content))])
    good = AsyncMock(return_value=resp(json.dumps({"role": "куратор", "title": "кеш жауап беру", "employee_name": "Арай"})))
    monkeypatch.setattr(ai.client.chat.completions, "create", good)
    assert run(REAL_CLASSIFY("текст")) == {"role": "Куратор", "title": "Кеш жауап беру", "employee_name": "Арай"}
    junk = AsyncMock(return_value=resp(json.dumps({"role": "жоқ тізімде", "title": "бірдеңе", "employee_name": "null"})))
    monkeypatch.setattr(ai.client.chat.completions, "create", junk)
    assert run(REAL_CLASSIFY("текст")) == {"role": "-", "title": "-", "employee_name": None}
    monkeypatch.setattr(ai.client.chat.completions, "create", AsyncMock(side_effect=RuntimeError("down")))
    assert run(REAL_CLASSIFY("текст"))["role"] == "-"


def test_ai_failure_does_not_block_complaint(env, monkeypatch):
    from bot import ai
    monkeypatch.setattr(service.ai, "classify", REAL_CLASSIFY)
    monkeypatch.setattr(ai.client.chat.completions, "create", AsyncMock(side_effect=RuntimeError("down")))
    post(env, TECH(), SMART, minimal_text())
    p = json.loads(C()[0]["parsed"])
    assert C()[0]["status"] == "open" and p["employee_role"] == "-" and p["issue_title"] == "-"


# ===================================================================== /pinall
class PinBot(FakeBot):
    def __init__(self, admin=True, can_pin=True):
        super().__init__()
        self.admin, self.can_pin, self.pinned = admin, can_pin, []

    async def get_chat_member(self, chat_id, user_id):
        return NS(status="administrator" if self.admin else "member")

    async def pin_chat_message(self, chat_id, msg_id, **kw):
        if not self.can_pin:
            raise RuntimeError("no rights")
        self.pinned.append(msg_id)


def _cmd_msg(bot, where):
    sent = []

    async def answer(text, **kw):
        bot._id += 1
        sent.append(text)
        return NS(message_id=bot._id, chat=NS(id=where[0]))
    return NS(chat=NS(id=where[0]), message_thread_id=where[1], is_topic_message=True,
              from_user=NS(id=99), answer=answer), sent


def test_pinall_pins_one_message_with_both_templates_in_smart_topic():
    b = PinBot()
    msg, sent = _cmd_msg(b, SMART)
    run(bm.cmd_pinall(msg, b))
    assert len(b.pinned) == 1 and len(sent) == 1                       # БІР pin хабарлама
    assert "Жаңа шағым" in sent[0] and "Шешім" in sent[0]               # екі шаблон бірге
    assert sent[0].count("<code>") == 2


def test_pinall_in_external_topic_pins_only_complaint_template():
    b = PinBot()
    msg, sent = _cmd_msg(b, EXT)
    run(bm.cmd_pinall(msg, b))
    assert len(b.pinned) == 1 and "Шағым шаблоны" in sent[0] and "Шешім" not in sent[0]


def test_pinall_ignored_for_non_admin_and_reports_missing_pin_right():
    b = PinBot(admin=False)
    msg, sent = _cmd_msg(b, SMART)
    run(bm.cmd_pinall(msg, b))
    assert sent == [] and b.pinned == []
    b = PinBot(can_pin=False)
    msg, sent = _cmd_msg(b, SMART)
    run(bm.cmd_pinall(msg, b))
    assert "Pin messages" in sent[-1]


# ===================================================================== шын чаттан шыққан ақаулар
def test_reminder_has_student_snippet_deadline_and_people(env):
    post(env, TECH(), SMART, complaint_text(student="Жансая Абил", desc="Куратор жауап бермейді, ата-ана ашулы"))
    env.clock.add(hours=1); run(service.flush(env.bot))
    t = env.bot.sent[-1].text
    for part in ("Еске салу", "#", "1 сағ", "Жансая Абил", "Куратор жауап бермейді", "@zhhaniya", "Кеш жауап беру",
                 "дедлайн", "Жауап күтілуде"):
        assert part in t, (part, t)


def test_digest_lines_include_student_names(env):
    for i, name in enumerate(("Айдос Нұрлан", "Мерей Асан")):
        post(env, TECH(), SMART, complaint_text(student=name, desc=f"мәселе {i}"))
    env.clock.add(hours=1); run(service.flush(env.bot))
    t = env.bot.sent[-1].text
    assert "Айдос Нұрлан" in t and "Мерей Асан" in t


class CbMsg:
    def __init__(self, bot, where, mid):
        self.chat, self.message_id, self._bot, self.where = NS(id=where[0]), mid, bot, where
        self.answered = []

    async def answer(self, text, **kw):
        self._bot._id += 1
        self.answered.append(self._bot._id)
        return NS(message_id=self._bot._id, chat=self.chat)


def test_reply_to_template_button_message_resolves_the_complaint(env):
    """Координатор «Шешім шаблоны» батырмасын басып, шыққан шаблонға reply жасайды (шын чаттағы ақау)."""
    m = post(env, TECH(), SMART, complaint_text())
    notify = env.bot.sent[0]
    cb_msg = CbMsg(env.bot, SMART, notify.id)
    cb = NS(message=cb_msg, answer=AsyncMock())
    run(bm.cb_template(cb))
    template_msg_id = cb_msg.answered[0]
    post(env, COORD(), SMART, resolution_text(), reply=template_msg_id)
    assert C()[0]["status"] == "resolved"


def test_reply_to_unknown_message_falls_back_to_name_matching(env):
    post(env, TECH(), SMART, complaint_text())
    post(env, COORD(), SMART, resolution_text(), reply=999999)     # шаблоны өшіп қалған/белгісіз хабарлама
    assert C()[0]["status"] == "resolved"


def test_progress_reply_marks_complaint_and_reminder_shows_it(env):
    post(env, TECH(), SMART, complaint_text())
    notify = env.bot.sent[0]
    n_react = len(env.bot.reactions)
    post(env, COORD(), SMART, "процессте, шешіліп жатыр", reply=notify.id)
    c = C()[0]
    assert c["status"] == "open" and c["progress_at"] and "Аружан" in c["progress_by"]
    assert len(env.bot.reactions) == n_react + 1 and env.bot.reactions[-1][2][0].emoji == "✍"
    env.clock.add(hours=1); run(service.flush(env.bot))
    assert "Процессте" in env.bot.sent[-1].text and "Аружан" in env.bot.sent[-1].text


def test_progress_variants_and_noise():
    for yes in ("процессте", "Шешіліп жатыр", "қарап жатырмыз", "в работе", "жұмыс істеп жатырмыз, бүгін шешеміз"):
        assert service.is_progress_note(yes), yes
    for no in ("рахмет", "ок", "Қайырлы күн", "түсінікті"):
        assert not service.is_progress_note(no), no


def test_progress_from_author_or_staff_is_ignored(env):
    tech = TECH()
    post(env, tech, SMART, complaint_text())
    notify = env.bot.sent[0]
    post(env, tech, SMART, "процессте", reply=notify.id)
    assert C()[0]["progress_at"] is None


def test_progress_then_template_still_resolves(env):
    post(env, TECH(), SMART, complaint_text())
    notify = env.bot.sent[0]
    post(env, COORD(), SMART, "процессте", reply=notify.id)
    post(env, COORD(), SMART, resolution_text(), reply=notify.id)
    assert C()[0]["status"] == "resolved"


# ===================================================================== шаблонсыз, бірақ түсінікті жазылған шағым
FREEFORM = """1)Айтжанова Айлин,
2)87054678221
2)ГЕО-21
3)Азатқан Толғанай
4)куратором өзіме ұнамайды жеке себептермен
@Meyrambekovich"""
EXTRACTED = {"is_complaint": True, "student_name": "Айтжанова Айлин", "phone": "87054678221",
             "subject": "ГЕО-21", "employee_name": "Азатқан Толғанай",
             "description": "куратором өзіме ұнамайды жеке себептермен"}


def test_freeform_numbered_message_is_accepted_automatically(env, monkeypatch):
    monkeypatch.setattr(service.ai, "extract_complaint", AsyncMock(return_value=EXTRACTED))
    post(env, TECH(), SMART, FREEFORM)
    c = C()
    assert len(c) == 1 and c[0]["status"] == "open"
    p = json.loads(c[0]["parsed"])
    assert p["student_name"] == "Айтжанова Айлин" and p["pairs"] == [["ГЕО", "21"]] and p["auto"] is True
    t = env.bot.sent[0].text
    assert "Айтжанова Айлин" in t and "@takogeo" in t and "@Meyrambekovich" in t        # 21 ағын ГЕО + супервайзер
    assert not any("Түзетіңіз" in m for m in env.bot.texts())


def test_freeform_chat_that_is_not_a_complaint_is_silent(env, monkeypatch):
    monkeypatch.setattr(service.ai, "extract_complaint", AsyncMock(return_value={
        "is_complaint": False, "student_name": None, "phone": None, "subject": None,
        "employee_name": None, "description": None}))
    post(env, TECH(), SMART, "Қайырлы күн!\nБүгін сабақ қай уақытта басталады екен?\nКім біледі, айтыңыздаршы")
    assert C() == [] and env.bot.sent == []


def test_freeform_missing_subject_gets_short_error_not_registered(env, monkeypatch):
    ext = dict(EXTRACTED, subject=None)
    monkeypatch.setattr(service.ai, "extract_complaint", AsyncMock(return_value=ext))
    post(env, TECH(), SMART, FREEFORM.replace("2)ГЕО-21", "2)пәні жоқ"))              # мәтінде пән коды мүлде жоқ
    assert C() == [] and "Пән" in env.bot.sent[0].text and len(env.bot.sent[0].text) < 200


def test_labelled_but_incomplete_message_is_completed_by_ai(env, monkeypatch):
    monkeypatch.setattr(service.ai, "extract_complaint", AsyncMock(return_value=EXTRACTED))
    partial = "Оқушы аты-жөні: Айтжанова Айлин\nПәні және потогы: ГЕО-21\nкуратором өзіме ұнамайды жеке себептермен деп жазды"
    post(env, TECH(), SMART, partial)
    assert len(C()) == 1 and json.loads(C()[0]["parsed"])["description"].startswith("куратором")


def test_freeform_fixed_by_edit_registers_once(env, monkeypatch):
    monkeypatch.setattr(service.ai, "extract_complaint", AsyncMock(return_value=dict(EXTRACTED, subject=None)))
    msg = post(env, TECH(), SMART, FREEFORM.replace("2)ГЕО-21", "2)пәні жоқ"))
    assert C() == []
    monkeypatch.setattr(service.ai, "extract_complaint", AsyncMock(return_value=EXTRACTED))
    msg.text = FREEFORM                                                              # менеджер пәнді қосып өңдеді (edit)
    run(bm.handle(msg, env.bot, True))
    assert len(C()) == 1 and C()[0]["status"] == "open"


def test_template_without_ai_still_works_and_short_chat_never_calls_ai(env, monkeypatch):
    spy = AsyncMock(return_value=None)
    monkeypatch.setattr(service.ai, "extract_complaint", spy)
    post(env, TECH(), SMART, complaint_text())            # дұрыс шаблон: OpenAI керек емес
    post(env, TECH(), SMART, "рахмет")
    spy.assert_not_awaited()


def test_progress_reply_to_the_reminder_message_itself(env):
    """Шын чаттағы жағдай: координатор еске салуға (reminder) «процессте, шешіліп жатыр» деп reply жасайды."""
    post(env, TECH(), SMART, complaint_text())
    env.clock.add(hours=1); run(service.flush(env.bot))
    reminder = env.bot.sent[-1]
    assert "Еске салу" in reminder.text
    coord = U(5, "Оразова Анеля | дүниежүзі тарихы координаторы", "aaneleka")
    post(env, coord, SMART, "процессте, шешіліп жатыр", reply=reminder.id)
    c = C()[0]
    assert c["status"] == "open" and c["progress_at"] and c["progress_by"] == "Оразова Анеля"
    assert env.bot.reactions[-1][2][0].emoji == "✍"
    env.clock.add(hours=1); run(service.flush(env.bot))
    assert "Процессте" in env.bot.sent[-1].text and "Оразова Анеля" in env.bot.sent[-1].text


# ===================================================================== /open
def test_open_list_is_detailed_and_does_not_ping(env):
    post(env, TECH(), SMART, complaint_text(student="Жансая Абил", subject="SMART ДЖТ-31"))
    post(env, EXTSELL(), EXT, complaint_text(student="Айдос Нұрлан", subject="SMART ГЕО-21", desc="басқа"))
    notify = env.bot.sent[0]
    post(env, COORD(), SMART, "процессте", reply=notify.id)
    env.clock.add(hours=2)
    text = "\n".join(service.fmt_open_list())
    for part in ("Ашық шағымдар (2)", "Жансая Абил", "Айдос Нұрлан", "ДЖТ-31", "Кеш жауап беру", "Есболқызы Арай",
                 "2 сағ", "дедлайн", "процессте", "Бекарыс Алишер", "Айтикова Аружан"):
        assert part in text, part
    assert "@" not in text                                   # атаулар @-сыз: қоңырау шалмайды
    assert 'href="https://t.me/c/' in text or "t.me/c/" in text


def test_open_list_flags_overdue_and_pending_states(env):
    post(env, TECH(), SMART, complaint_text())
    env.clock.add(hours=30)
    assert "дедлайн өтті" in service.fmt_open_list()[0] and "🚨" in service.fmt_open_list()[0]
    env.clock.set(2026, 10, 17, 22, 0)                      # сенбі түн: жаңа шағым да бірден қабылданады
    post(env, TECH(), SMART, complaint_text(student="Түнгі Оқушы", desc="түнде"))
    assert "Түнгі Оқушы" in "\n".join(service.fmt_open_list())


def test_open_list_is_split_into_telegram_sized_chunks(env):
    for i in range(40):
        post(env, TECH(), SMART, complaint_text(student=f"Оқушы Нөмір{i}", desc=f"мәселе {i} " + "көп мәтін " * 5))
    chunks = service.fmt_open_list()
    assert len(chunks) >= 2 and all(len(c) < 4096 for c in chunks)
    assert sum(c.count("Оқушы Нөмір") for c in chunks) == 40


def test_open_in_group_shows_only_that_topic_and_expires(env):
    post(env, TECH(), SMART, complaint_text(student="Смарт Оқушы"))
    post(env, TECH(), JUNIOR, complaint_text(student="Джуниор Оқушы", subject="JUNIOR БИО", desc="басқа"))
    out = []

    async def answer(text, **kw):
        env.bot._id += 1
        out.append(text)
        return NS(message_id=env.bot._id, chat=NS(id=SMART[0]))
    msg = NS(chat=NS(id=SMART[0], type="supergroup"), message_thread_id=SMART[1], is_topic_message=True,
             from_user=NS(id=1), answer=answer)
    run(bm.cmd_open(msg, env.bot))
    assert "Смарт Оқушы" in out[0] and "Джуниор Оқушы" not in out[0]
    assert db.q1("SELECT COUNT(*) n FROM ephemeral")["n"] == 1


def test_open_in_private_needs_admin(env):
    class AdminBot(FakeBot):
        def __init__(self, admin):
            super().__init__()
            self.admin = admin

        async def get_chat_member(self, chat_id, user_id):
            return NS(status="administrator" if self.admin else "member")
    post(env, TECH(), SMART, complaint_text())
    for admin, expected in ((False, "админдеріне"), (True, "Ашық шағымдар")):
        out = []

        async def answer(text, **kw):
            out.append(text)
            return NS(message_id=1, chat=NS(id=1))
        msg = NS(chat=NS(id=1, type="private"), from_user=NS(id=9), answer=answer)
        run(bm.cmd_open(msg, AdminBot(admin)))
        assert expected in out[0]


# ===================================================================== «қабылданды / кб береміз» да процесс
def test_acknowledgement_phrases_count_as_progress():
    for yes in ("кеш жарық, қабылданды\nкб береміз", "қабылданды", "кб береміз", "Кері байланыс береміз",
                "принято", "сообщим позже"):
        assert service.is_progress_note(yes), yes


def test_ack_reply_from_real_chat_marks_progress_without_ai(env):
    post(env, TECH(), SMART, complaint_text())
    coord = U(6, "ГҮЛНАЗ СЕРІҚҚЫЗЫ | ГЕОГРАФИЯ КООРДИНАТОРЫ")
    db.x("INSERT OR REPLACE INTO users(user_id, username, full_name) VALUES(?,?,?)", 6, None, coord.full_name)
    post(env, coord, SMART, "кеш жарық, қабылданды\nкб береміз", reply=env.bot.sent[0].id)
    assert C()[0]["progress_at"] and C()[0]["status"] == "open"


def test_ai_fallback_detects_progress_in_unusual_wording(env, monkeypatch):
    post(env, TECH(), SMART, complaint_text())
    monkeypatch.setattr(service.ai, "classify_reply", AsyncMock(return_value={"kind": "progress"}))
    post(env, COORD(), SMART, "ертең таңертең оқушымен сөйлесіп көрем", reply=env.bot.sent[0].id)
    assert C()[0]["progress_at"]


def test_ai_fallback_not_used_for_short_noise_or_staff(env, monkeypatch):
    spy = AsyncMock(return_value={"kind": "progress"})
    monkeypatch.setattr(service.ai, "classify_reply", spy)
    tech = TECH()
    post(env, tech, SMART, complaint_text())
    post(env, tech, SMART, "ертең таңертең оқушымен сөйлесіп көрем", reply=env.bot.sent[0].id)   # жазған адам - есепке алынбайды
    post(env, COORD(), SMART, "ок", reply=env.bot.sent[0].id)                                    # тым қысқа
    spy.assert_not_awaited()
    assert C()[0]["progress_at"] is None


# ===================================================================== жаңа шағым (🔔) хабарламасына «в процессе»
@pytest.mark.parametrize("phrase", ["в процессе", "В процессе, решаем", "в работе", "Процессе", "процессте"])
def test_progress_reply_to_the_new_complaint_notification(env, phrase):
    post(env, TECH(), SMART, complaint_text())
    notify = env.bot.sent[0]
    assert "🔔" in notify.text
    post(env, COORD(), SMART, phrase, reply=notify.id)
    c = C()[0]
    assert c["status"] == "open" and c["progress_at"], phrase
    assert env.bot.reactions[-1][2][0].emoji == "✍"


def test_progress_reply_to_the_forwarded_external_complaint_and_its_notification(env):
    post(env, EXTSELL(), EXT, complaint_text())
    c = C()[0]
    notify = [m for m in env.bot.sent if "🔔" in m.text][0]
    post(env, COORD(), SMART, "в процессе", reply=notify.id)                 # 🔔 хабарламасына
    assert C()[0]["progress_at"]
    db.update("complaints", c["id"], progress_at=None)
    post(env, COORD(), SMART, "в процессе", reply=c["target_msg"])           # көшірілген шағымның өзіне
    assert C()[0]["progress_at"]


def test_program_question_repeats_hourly_at_any_time(env):
    env.clock.set(2026, 10, 18, 3, 0)                       # жексенбі түн - жұмыс уақытынан тыс
    post(env, EXTSELL(), EXT, complaint_text(subject="ГЕО-31"))
    q = lambda: [m for m in env.bot.sent if "SMART па" in m.text]  # noqa: E731
    assert len(q()) == 1                                    # бірінші сұрақ бірден
    env.clock.add(minutes=30); run(service.flush(env.bot)); assert len(q()) == 1
    env.clock.add(minutes=31); run(service.flush(env.bot)); assert len(q()) == 2    # 1 сағаттан кейін қайталанады
    env.clock.add(hours=1); run(service.flush(env.bot)); assert len(q()) == 3


# ===================================================================== /close
class CloseBot(FakeBot):
    async def get_chat_member(self, chat_id, user_id):
        return NS(status="administrator" if user_id == 99 else "member")


def _close(env, args, reply_to=None, admin=True):
    bot = env.bot
    out = []

    async def reply(text, **kw):
        out.append(text)
    msg = NS(chat=NS(id=SMART[0]), from_user=NS(id=99 if admin else 5, full_name="Админ Аты"), reply=reply,
             reply_to_message=NS(message_id=reply_to) if reply_to else None)
    run(bm.cmd_close(msg, NS(args=args), bot))
    return out


@pytest.fixture
def cbot(env):
    env.bot = env.bot.__class__ and CloseBot()
    return env


def test_close_by_id_with_reason_writes_reason_into_the_excel_answer(cbot):
    env = cbot
    post(env, TECH(), SMART, complaint_text())
    out = _close(env, f"#{C()[0]['id']} оқушы курстан бас тартты, шағым қайтарылды")
    assert "жабылды" in out[0] and C()[0]["status"] == "resolved"
    assert json.loads(C()[0]["final"])["edu_response"].startswith("оқушы курстан")
    assert "шешілді" in env.bot.sent[-1].text


def test_close_as_reply_to_the_complaint_or_bot_message(cbot):
    env = cbot
    m = post(env, TECH(), SMART, complaint_text())
    out = _close(env, "", reply_to=env.bot.sent[0].id)           # 🔔 хабарламасына reply
    assert "жабылды" in out[0] and C()[0]["status"] == "resolved"
    m2 = post(env, TECH(), SMART, complaint_text(student="Басқа Бала", desc="басқа"))
    out = _close(env, "дубль", reply_to=m2.message_id)           # шағымның өзіне reply
    assert [c["status"] for c in C()] == ["resolved", "resolved"]


def test_close_requires_admin_and_reports_unknown_or_double_close(cbot):
    env = cbot
    post(env, TECH(), SMART, complaint_text())
    cid = C()[0]["id"]
    assert _close(env, str(cid), admin=False) == [] and C()[0]["status"] == "open"
    assert "Қолданылуы" in _close(env, "")[0]
    _close(env, str(cid))
    assert "Табылмады" in _close(env, str(cid))[0]
    assert "Табылмады" in _close(env, "9999")[0]


# ===================================================================== /delete (қате шағым)
def _delete(env, args, reply_to=None, admin=True):
    out = []

    async def reply(text, **kw):
        out.append(text)
        return NS(chat=NS(id=SMART[0]), message_id=777)

    async def delete():
        out.append("<cmd-deleted>")
    msg = NS(chat=NS(id=SMART[0]), from_user=NS(id=99 if admin else 5, full_name="Админ"), reply=reply,
             delete=delete, reply_to_message=NS(message_id=reply_to) if reply_to else None)
    run(bm.cmd_delete(msg, NS(args=args), env.bot))
    return out


def test_delete_keeps_bot_messages_visible_as_a_tombstone_and_removes_the_complaint(cbot):
    env = cbot
    post(env, TECH(), SMART, complaint_text())
    env.clock.add(hours=1); run(service.flush(env.bot))                 # еске салу да болсын
    cid = C()[0]["id"]
    bot_msg_ids = {m.id for m in env.bot.sent}
    out = _delete(env, str(cid))
    assert "өшірілді" in out[0]
    assert C()[0]["status"] == "deleted"
    assert env.bot.deleted == []                                        # чаттан ештеңе өшпейді (басқалар көре алады)
    assert {mid for _, mid, _ in env.bot.edited} >= bot_msg_ids         # 🔔 мен еске салу «🗑 өшірілді» белгісіне ауысты
    assert all("қате шағым ретінде өшірілді" in t for _, _, t in env.bot.edited)
    assert "Ашық шағым жоқ" in service.fmt_open_list()[0]
    n = len(env.bot.sent)
    env.clock.add(hours=3); run(service.flush(env.bot))
    assert len(env.bot.sent) == n                                      # енді еске салу жоқ
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(exporter.build(include_open=True)))
    assert all(ws.max_row == 1 for ws in wb)                           # тек тақырып қатары, дерек жоқ


def test_delete_by_reply_and_works_for_resolved_complaints_too(cbot):
    env = cbot
    m = post(env, TECH(), SMART, complaint_text())
    post(env, COORD(), SMART, resolution_text(), reply=m.message_id)
    assert C()[0]["status"] == "resolved"
    _delete(env, "", reply_to=m.message_id)
    assert C()[0]["status"] == "deleted"
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(exporter.build()))
    assert all(ws.max_row == 1 for ws in wb)                           # Excel-ге түспейді


def test_delete_requires_admin_and_handles_unknown_ids(cbot):
    env = cbot
    post(env, TECH(), SMART, complaint_text())
    cid = C()[0]["id"]
    assert _delete(env, str(cid), admin=False) == [] and C()[0]["status"] == "open"
    assert "Қолданылуы" in _delete(env, "")[0]
    assert "Табылмады" in _delete(env, "9999")[0]
    _delete(env, str(cid))
    assert "Табылмады" in _delete(env, str(cid))[0]


def test_deleted_external_complaint_ignores_later_replies_and_deletes_nothing(cbot):
    env = cbot
    post(env, EXTSELL(), EXT, complaint_text())
    c = C()[0]
    copied_id = c["target_msg"]
    _delete(env, str(c["id"]))
    assert (SMART[0], copied_id) not in env.bot.deleted                # көшірме де өшпейді
    post(env, COORD(), SMART, resolution_text(), reply=copied_id)
    assert C()[0]["status"] == "deleted"                               # өшірілген шағым қайта жабылмайды


# ===================================================================== шаблонсыз жазылған ШЕШІМ (процесс емес)
FREE_RESOLUTION = ("Жауапты коорд оқушымен байланысты, кураторын рейтинг шыққан соң ауыстырып береді. "
                   "Оқушы шешімге риза")


def _ai_resolution(**kw):
    base = {"kind": "resolution", "solution": "Оқушымен байланысты, куратор рейтинг шыққан соң ауыстырылады",
            "satisfied": None, "justified": None, "stayed": None}
    base.update(kw)
    return AsyncMock(return_value=base)


class CbMsgDraft:
    def __init__(self, chat_id, mid):
        self.chat, self.message_id = NS(id=chat_id), mid


def tap(env, cid, field, label_idx, user=None):
    cb = NS(data=f"dr:{cid}:{field}:{label_idx}", from_user=user or COORD(), message=CbMsgDraft(SMART[0], 1),
            answer=AsyncMock())
    run(bm.cb_draft(cb, env.bot))
    return cb


def test_free_text_resolution_is_not_progress_and_asks_only_what_is_missing_with_buttons(env, monkeypatch):
    monkeypatch.setattr(service.ai, "classify_reply", _ai_resolution(satisfied="Иә"))
    post(env, TECH(), SMART, complaint_text(student="Айтжанова Айлин"))
    n_react = len(env.bot.reactions)
    post(env, COORD(), SMART, FREE_RESOLUTION, reply=env.bot.sent[0].id)
    c = C()[0]
    assert c["status"] == "open" and c["progress_at"] is None            # «процессте» деп белгіленбейді
    assert len(env.bot.reactions) == n_react + 1 and env.bot.reactions[-1][2][0].emoji == "👌"   # ✍ емес, 👌
    q = env.bot.sent[-1]
    assert "негізделген" in q.text and "риза" not in q.text              # риза ма - мәтіннен алынды, сұралмайды
    buttons = [b.text for r in q.kb.inline_keyboard for b in r]
    assert buttons == ["Негізделген", "Жартылай негізделген", "Негізделмеген", "КБ бермеді"]


def test_buttons_complete_the_resolution_and_close_the_complaint(env, monkeypatch):
    monkeypatch.setattr(service.ai, "classify_reply", _ai_resolution(satisfied="Иә"))
    post(env, TECH(), SMART, complaint_text(student="Айтжанова Айлин"))
    post(env, COORD(), SMART, FREE_RESOLUTION, reply=env.bot.sent[0].id)
    cid = C()[0]["id"]
    tap(env, cid, "justified", 0)                                         # Негізделген
    assert C()[0]["status"] == "open" and env.bot.edited[-1][2].count("курста") == 1     # келесі сұрақ сол хабарлама ішінде
    tap(env, cid, "stayed", 0)                                            # Иә
    c = C()[0]
    f = json.loads(c["final"])
    assert c["status"] == "resolved" and (f["justified"], f["satisfied"], f["stayed"]) == ("Негізделген", "Иә", "Иә")
    assert "шешім толық" in env.bot.edited[-1][2]


def test_only_a_resolver_can_tap_the_buttons(env, monkeypatch):
    monkeypatch.setattr(service.ai, "classify_reply", _ai_resolution(satisfied="Иә"))
    tech = TECH()
    post(env, tech, SMART, complaint_text())
    post(env, COORD(), SMART, FREE_RESOLUTION, reply=env.bot.sent[0].id)
    cb = tap(env, C()[0]["id"], "justified", 0, user=tech)               # шағымды жазған қызметкер
    cb.answer.assert_awaited_with("Тек жауапты координатор/супервайзер", show_alert=True)
    assert "justified" not in json.loads(C()[0]["draft"])


def test_short_followup_reply_and_unmapped_reply_continue_the_same_draft(env, monkeypatch):
    """Шын чат: алдымен шешім, содан кейін «Оқушы риза, курста қалды» (өшіп қалған хабарламаға reply)."""
    monkeypatch.setattr(service.ai, "classify_reply", AsyncMock(side_effect=[
        {"kind": "resolution", "solution": "Жеке сөйлесу жүргізілді, кураторға ескертілді және куратор ауыстырылды",
         "satisfied": None, "justified": "Негізделген", "stayed": None},
        {"kind": "other", "solution": None, "satisfied": "Иә", "justified": None, "stayed": "Иә"}]))
    post(env, TECH(), SMART, complaint_text(student="Жәнібекқызы Аяжан"))
    coord = COORD()
    post(env, coord, SMART, "Жәнібекқызы Аяжамен жеке сөйлесу жүргізілді, дәлелдер қабылданды. Кураторға ескертілді "
                            "және оқушының кураторын ауыстырдық", reply=env.bot.sent[0].id)
    assert C()[0]["status"] == "open" and "риза" in env.bot.sent[-1].text      # негізделген мәтіннен алынды -> риза ма сұралады
    post(env, coord, SMART, "Оқушы риза, курста қалды", reply=999999)         # өшіп қалған/белгісіз хабарламаға reply
    c = C()[0]
    f = json.loads(c["final"])
    assert c["status"] == "resolved" and (f["justified"], f["satisfied"], f["stayed"]) == ("Негізделген", "Иә", "Иә")


def test_followup_without_any_reply_works_for_the_same_coordinator_only(env, monkeypatch):
    monkeypatch.setattr(service.ai, "classify_reply", AsyncMock(side_effect=[
        {"kind": "resolution", "solution": "Оқушымен сөйлестік, куратор ауыстырылды", "satisfied": None,
         "justified": "Негізделген", "stayed": None},
        {"kind": "other", "solution": None, "satisfied": "Иә", "justified": None, "stayed": "Иә"}]))
    post(env, TECH(), SMART, complaint_text())
    coord = COORD()
    post(env, coord, SMART, "Оқушымен сөйлестік, куратор ауыстырылды, шағым дұрыс болды", reply=env.bot.sent[0].id)
    other = U(77, "Басқа Координатор", "somebody77")
    db.x("INSERT OR REPLACE INTO users(user_id, username, full_name) VALUES(?,?,?)", 77, "somebody77", other.full_name)
    post(env, other, SMART, "Оқушы риза, курста қалды")                         # басқа адам: жоба жалғаспайды
    assert C()[0]["status"] == "open"
    post(env, coord, SMART, "Оқушы риза, курста қалды")                          # reply-сыз, сол координатор
    assert C()[0]["status"] == "resolved"


def test_draft_expires_after_six_hours(env, monkeypatch):
    monkeypatch.setattr(service.ai, "classify_reply", _ai_resolution(satisfied="Иә"))
    post(env, TECH(), SMART, complaint_text())
    coord = COORD()
    post(env, coord, SMART, FREE_RESOLUTION, reply=env.bot.sent[0].id)
    assert service.active_draft_complaint(coord.id, SMART[0], SMART[1]) is not None
    env.clock.add(hours=7)
    assert service.active_draft_complaint(coord.id, SMART[0], SMART[1]) is None


def test_classify_reply_failure_falls_back_to_keywords_only(env, monkeypatch):
    monkeypatch.setattr(service.ai, "classify_reply", AsyncMock(return_value=None))
    post(env, TECH(), SMART, complaint_text())
    post(env, COORD(), SMART, FREE_RESOLUTION, reply=env.bot.sent[0].id)      # AI жоқ, сөздер жоқ -> ештеңе істемейді
    assert C()[0]["progress_at"] is None and C()[0]["status"] == "open"
    post(env, COORD(), SMART, "қабылданды", reply=env.bot.sent[0].id)           # сөздер бойынша процесс
    assert C()[0]["progress_at"]


def test_roster_chem_11_coordinator_replaced():
    from bot.roster import responsible_for
    coords, sups, _ = responsible_for("smart", ["ХИМ"], "11")
    usernames = [p.username for p in coords]
    assert usernames == ["beckhac_e", "yergalieva5"]
    assert all("Жеткіншек" not in p.name for p in coords)


def test_roster_updated_usernames_and_name_only_people():
    from types import SimpleNamespace as N
    from bot import people
    from bot.roster import ALL_PEOPLE, responsible_for
    names = {p.username.lower() for p in ALL_PEOPLE if p.username}
    for old in ("aitoreshkaaiz", "amikowm", "guldanari", "marishka_berekeeva", "mkhmdl", "oralbayeva_aruzhan", "yerganatt",
                "zere140qazrus", "parvinasavronova", "qaselok", "talshynkaiyrgazina", "tannur_eng", "altynay_kn"):
        assert old not in names, old
    for new in ("aitoreshka_juz40quqyq", "amina_info140", "guldana_mc", "mariyam_ruslit140", "mereyka_bio140",
                "oralbayeva_tarikhjuz40", "yerganat_dzht140", "zere_adebtil140", "parvina_bio140", "qaselok_bio140",
                "toshageom140", "mereiphys", "assema_chemistry140", "dikosh_juz40"):
        assert new in names, new
    for stream in ("01", "11"):                                     # АНГЛ координаторы Сәндібек Мадина
        coords, _, _ = responsible_for("smart", ["АНГЛ"], stream)
        assert [p.username for p in coords] == ["sandibekmadina"]
    coords, sups, _ = responsible_for("smart", ["ФИЗ"], "21")       # ник жоқ координатор аты бойынша танылады
    assert any(p.name == "Каненова Алтынай" for p in coords)
    assert people.is_roster_user(N(username=None, full_name="Алтынай Каненова | Физика КООРДИНАТОР"))


def test_long_concrete_solution_is_accepted_without_ai_gate(env, monkeypatch):
    """Шын чаттағы жағдай: толық шешім «нақты емес» деп қайтарылған."""
    gate = AsyncMock(return_value={"ok": False, "reason": "қатты талап"})
    monkeypatch.setattr(service.ai, "check_solution", gate)
    post(env, TECH(), SMART, complaint_text(student="Жибек Тныштыбай"))
    text = ("Оқушы аты-жөні: Жибек Тныштыбай\nМәселенің шешімі: оқушы тақырыптардан қалып қойған, сол үшін қыркүйек ағынына "
            "ауыстыдық, сол қалып қойған тақырыптарынан керi бастайды\nМәселе шындыққа негізделген бе? Негізделген\n"
            "Оқушы риза ма? Иә, риза\nОқушы курста қалды ма? Иә")
    post(env, COORD(), SMART, text, reply=env.bot.sent[0].id)
    assert C()[0]["status"] == "resolved"
    gate.assert_not_awaited()
    f = json.loads(C()[0]["final"])
    assert f["satisfied"] == "Иә" and f["stayed"] == "Иә"                  # «Иә, риза» да «Иә» деп танылды


def test_short_vague_solution_still_goes_through_the_ai_gate(env, monkeypatch):
    gate = AsyncMock(return_value={"ok": False, "reason": "нақтылаңыз"})
    monkeypatch.setattr(service.ai, "check_solution", gate)
    post(env, TECH(), SMART, complaint_text())
    short = resolution_text().replace("Оқушымен байланысқа шықтық, себебі анықталды, кураторға ескерту жасалды.",
                                      "Мәселе шешілді, оқушымен сөйлестім, бәрі жақсы")
    post(env, COORD(), SMART, short, reply=env.bot.sent[0].id)
    gate.assert_awaited()
    assert C()[0]["status"] == "open"


# ===================================================================== ағын ай аты бойынша (шілде 01 ... қараша 41)
@pytest.mark.parametrize("label,program,pairs", [
    ("География SMART ҚЫРКҮЙЕК", "smart", [("ГЕО", "21")]),
    ("ХИМ шілде", None, [("ХИМ", "01")]),
    ("Физика, математика тамыз", None, [("ФИЗ", "11"), ("МАТ", "11")]),
    ("SMART БИО қараша", "smart", [("БИО", "41")]),
    ("Сентябрь ГЕО СМАРТ", "smart", [("ГЕО", "21")]),
    ("ДЖТ қазан потогы", None, [("ДЖТ", "31")]),
    ("SMART ГЕО-11 қыркүйек", "smart", [("ГЕО", "11")]),                 # нақты нөмір ай атынан басым
    ("қыркүйектегі география", None, [("ГЕО", "21")]),
    ("SMART тарихы шілдедегі", "smart", [("ТАРИХ", "01")]),
])
def test_stream_from_month_names(label, program, pairs):
    p, got = T.parse_subject(label)
    assert p == program and got == pairs


def test_stream_41_is_known_but_has_no_coordinators_yet():
    from bot.roster import STREAM_NAMES, responsible_for
    assert STREAM_NAMES["41"] == "41 қараша"
    coords, sups, notes = responsible_for("smart", ["ГЕО"], "41")
    assert coords == [] and [p.username for p in sups] == ["Meyrambekovich"]


FREEFORM_MONTH = """Қайырлы күн
Биназир Төлеген
+77718986439
tolegen.binazir@icloud.com
География SMART ҚЫРКҮЙЕК
Аякөз Қыдырбай кураторы
Мәселе: кураторыммен био кураторыммен сөйлескендей ашылып сөйлесе алмаймын, геодан қиналамын дейді
Байланыса аласыздар ма?
@nurzhannn_geo"""


def test_screenshot_case_month_name_gives_stream_automatically(env, monkeypatch):
    monkeypatch.setattr(service.ai, "extract_complaint", AsyncMock(return_value={
        "is_complaint": True, "student_name": "Биназир Төлеген", "phone": "+77718986439",
        "subject": "География SMART ҚЫРКҮЙЕК", "employee_name": "Аякөз Қыдырбай",
        "description": "кураторыммен сөйлесе алмаймын, геодан қиналамын дейді"}))
    post(env, TECH(), SMART, FREEFORM_MONTH)
    assert len(C()) == 1 and C()[0]["status"] == "open"
    p = json.loads(C()[0]["parsed"])
    assert p["pairs"] == [["ГЕО", "21"]] and p["subject_label"] == "SMART ГЕО-21"          # Excel форматы
    t = env.bot.sent[0].text
    assert "@takogeo" in t and "@nurzhannn_geo" in t and "@Meyrambekovich" in t and not any("Түзетіңіз" in m for m in env.bot.texts())


def test_subject_label_is_normalised_for_excel(env):
    post(env, TECH(), SMART, complaint_text(subject="Физика, математика шілде"))
    assert json.loads(C()[0]["parsed"])["subject_label"] == "SMART ФИЗ-01, МАТ-01"
    post(env, TECH(), JUNIOR, complaint_text(student="Басқа Бала", subject="Биология", desc="x"))
    assert json.loads(C()[1]["parsed"])["subject_label"] == "JUNIOR БИО"


def test_junior_topic_does_not_require_stream_even_without_the_word_junior(env):
    post(env, TECH(), JUNIOR, complaint_text(subject="Биология"))
    assert len(C()) == 1 and C()[0]["program"] == "junior"
    post(env, TECH(), SMART, complaint_text(student="Басқа Бала", subject="Биология", desc="x"))   # Smart-та ағын керек
    assert len(C()) == 2 and "Ағын көрсетілмеген" in env.bot.sent[-1].text      # Smart-та ағынсыз да қабылданады


# ===================================================================== ортақ еске салу тізіміне (digest) reply
def _two_open_with_digest(env):
    post(env, TECH(), SMART, complaint_text(student="Жансая Абил", desc="бірінші мәселе"))
    post(env, TECH(), SMART, complaint_text(student="Шұғыла Төлеген", desc="екінші мәселе"))
    env.clock.add(hours=1); run(service.flush(env.bot))
    digest = env.bot.sent[-1]
    assert "Шешілмеген шағымдар (2)" in digest.text
    return digest


def test_progress_reply_to_digest_is_attributed_by_student_name(env, monkeypatch):
    monkeypatch.setattr(service.ai, "classify_reply", AsyncMock(return_value={"kind": "progress"}))   # ұзын мәтін: OpenAI шешеді
    digest = _two_open_with_digest(env)
    n_react = len(env.bot.reactions)
    text = ("Шұғыла Төлеген- кураторы Сагидуллиева Аида. Оқушы осы уақытқа дейін жақсы оқып келіпті.\n"
            "Бүгін кешке кураторы жеке митта сөйлеседі, жауабын жазатын боламын")
    post(env, COORD(), SMART, text, reply=digest.id)
    by_student = {json.loads(c["parsed"])["student_name"]: c for c in C()}
    assert by_student["Шұғыла Төлеген"]["progress_at"] and by_student["Жансая Абил"]["progress_at"] is None
    assert len(env.bot.reactions) == n_react + 1 and env.bot.reactions[-1][2][0].emoji == "✍"


def test_digest_reply_with_complaint_number_or_single_candidate(env):
    digest = _two_open_with_digest(env)
    first_id = C()[0]["id"]
    post(env, COORD(), SMART, f"#{first_id} қабылданды, кб береміз", reply=digest.id)
    assert C()[0]["progress_at"] and C()[1]["progress_at"] is None
    for c in C():                                           # тізімде бір ғана шағым болса - нөмір/аты керек емес
        db.update("complaints", c["id"], status="deleted")
    post(env, TECH(), SMART, complaint_text(student="Жалғыз Оқушы", desc="жалғыз"))
    env.clock.add(hours=2); run(service.flush(env.bot))
    single = env.bot.sent[-1]
    post(env, COORD(), SMART, "қабылданды, кб береміз", reply=single.id)
    assert C()[-1]["progress_at"]


def test_ambiguous_digest_reply_asks_which_complaint_and_marks_nothing(env):
    digest = _two_open_with_digest(env)
    post(env, COORD(), SMART, "қабылданды, кб береміз", reply=digest.id)
    assert all(c["progress_at"] is None for c in C())
    assert "Қай шағым" in env.bot.sent[-1].text


def test_digest_reply_with_resolution_template_closes_the_named_student(env):
    digest = _two_open_with_digest(env)
    post(env, COORD(), SMART, resolution_text(student="Шұғыла Төлеген"), reply=digest.id)
    by_student = {json.loads(c["parsed"])["student_name"]: c["status"] for c in C()}
    assert by_student == {"Жансая Абил": "open", "Шұғыла Төлеген": "resolved"}


# ===================================================================== скриншот + қысқа caption (сипаттама суретте)
SHOT_CAPTION = """Қайырлы кеш🙌
Жәнібекқызы Аяжан
+77770337269
Т-21
Көркем Оразғалиева
@oralbayeva_tarikhjuz40"""
NO_DESC = {"is_complaint": False, "student_name": "Жәнібекқызы Аяжан", "phone": "+77770337269", "subject": None,
           "employee_name": "Көркем Оразғалиева", "description": None}
OCR_TEXT = "Жәнібекқызы Аяжан 87770337269\nТарих пәні 21ші ағым\nКуратор ауыстыруым себебі жұмысын дұрыс атқара алмайды"


def photo_msg(env, caption):
    m = M(TECH(), SMART, caption=caption, ctype="photo")
    m.photo = [NS(file_id="small"), NS(file_id="big")]
    return m


def test_screenshot_with_structured_caption_is_accepted_and_description_read_from_the_picture(env, monkeypatch):
    full = dict(NO_DESC, is_complaint=True, subject=None,
                description="Куратор ауыстыруым себебі жұмысын дұрыс атқара алмайды")
    monkeypatch.setattr(service.ai, "extract_complaint", AsyncMock(side_effect=[NO_DESC, full]))
    shot = AsyncMock(return_value=OCR_TEXT)
    monkeypatch.setattr(service.ai, "read_screenshot", shot)
    run(bm.handle(photo_msg(env, SHOT_CAPTION), env.bot, False))
    shot.assert_awaited_once()
    c = C()
    assert len(c) == 1 and c[0]["status"] == "open"
    p = json.loads(c[0]["parsed"])
    assert p["description"].startswith("Куратор ауыстыруым") and p["pairs"] == [["ТАРИХ", "21"]]   # «Т-21» = ТАРИХ-21
    assert p["student_name"] == "Жәнібекқызы Аяжан" and p["subject_label"] == "SMART ТАРИХ-21"
    assert "@oralbayeva_tarikhjuz40" in env.bot.sent[0].text                                      # тарих супервайзері


def test_screenshot_ocr_unavailable_still_accepts_with_attachment_placeholder(env, monkeypatch):
    monkeypatch.setattr(service.ai, "extract_complaint", AsyncMock(return_value=NO_DESC))
    run(bm.handle(photo_msg(env, SHOT_CAPTION), env.bot, False))                                    # read_screenshot -> None
    c = C()
    assert len(c) == 1 and "тіркемеде" in json.loads(c[0]["parsed"])["description"]


def test_message_without_screenshot_and_without_description_asks_for_it(env, monkeypatch):
    ext = dict(NO_DESC, is_complaint=True)
    monkeypatch.setattr(service.ai, "extract_complaint", AsyncMock(return_value=ext))
    post(env, TECH(), SMART, SHOT_CAPTION)                                                          # сурет те, сипаттама да жоқ
    assert C() == [] and "сипаттамасы" in env.bot.sent[0].text.lower()


def test_chat_with_picture_but_no_student_is_still_ignored(env, monkeypatch):
    monkeypatch.setattr(service.ai, "extract_complaint", AsyncMock(return_value={
        "is_complaint": False, "student_name": None, "phone": None, "subject": None,
        "employee_name": None, "description": None}))
    run(bm.handle(photo_msg(env, "Міне кестемен таныс болыңыздар, ертең сабақ осылай болады"), env.bot, False))
    assert C() == []


@pytest.mark.parametrize("text,expected", [
    ("Т-21", "Т-21"), ("Көркем Оразғалиева\n+7 777 033 72 69\nГЕО-31 қазан", "ГЕО-31"), ("тарих - 11", "ТАРИХ-11"),
    ("87770337269", None), ("сәлем, қалайсыз", None), ("Асан-21 жас", None),
])
def test_find_subject_in_text(text, expected):
    assert T.find_subject_in_text(text) == expected


def test_openai_missed_subject_but_text_has_code_is_found_deterministically(env, monkeypatch):
    monkeypatch.setattr(service.ai, "extract_complaint", AsyncMock(return_value=dict(EXTRACTED, subject=None)))
    post(env, TECH(), SMART, FREEFORM)                                               # мәтінде «2)ГЕО-21» бар
    assert len(C()) == 1 and json.loads(C()[0]["parsed"])["pairs"] == [["ГЕО", "21"]]


# ===================================================================== боттың өшіруі журналға жазылады + /status
def test_every_bot_deletion_is_logged_and_status_shows_it(env, monkeypatch):
    post(env, TECH(), SMART, None, ctype="photo")                       # «мәтінсіз» ескерту (ttl 15 мин)
    assert db.q1("SELECT COUNT(*) n FROM deletions")["n"] == 0
    env.clock.add(minutes=16); run(service.flush(env.bot))              # TTL өтті -> бот өшірді
    d = db.q("SELECT * FROM deletions")
    assert len(d) == 1 and "TTL" in d[0]["reason"]
    status = service.fmt_status()
    assert "Боттың соңғы өшірулері" in status and "TTL" in status and "Шағымдар" in status


def test_important_messages_are_never_auto_deleted(env):
    """🔔, еске салу, ✅ шешілді: уақыт өткенде де чаттан өшпейді (тек қате/шаблон/ /open уақытша)."""
    m = post(env, TECH(), SMART, complaint_text())
    env.clock.add(hours=1); run(service.flush(env.bot))
    post(env, COORD(), SMART, resolution_text(), reply=env.bot.sent[0].id)
    env.clock.add(days=3); run(service.flush(env.bot))
    assert env.bot.deleted == []
    assert db.q1("SELECT COUNT(*) n FROM ephemeral")["n"] == 0
    assert db.q1("SELECT COUNT(*) n FROM deletions")["n"] == 0


def test_errors_are_recorded_for_status(env):
    import logging
    logging.getLogger("bot.test").error("сынақ қатесі")
    assert any("сынақ қатесі" in e["text"] for e in db.q("SELECT * FROM errors"))
    assert "сынақ қатесі" in service.fmt_status()


# ===================================================================== жауапқа бот реакциясы және «неге үндемеді?» журналы
def test_visible_reactions_for_each_outcome(env, monkeypatch):
    post(env, TECH(), SMART, complaint_text())
    notify = env.bot.sent[0]
    post(env, COORD(), SMART, "в процессе", reply=notify.id)
    assert env.bot.reactions[-1][2][0].emoji == "✍"
    post(env, COORD(), SMART, resolution_text(), reply=notify.id)
    assert env.bot.reactions[-1][2][0].emoji == "👍" and C()[0]["status"] == "resolved"


def test_long_free_text_resolution_with_a_photo_caption_gets_reaction_and_question(env, monkeypatch):
    """Шын чат: координатор скриншот + caption (шаблонсыз шешім) жіберді, бот үнсіз қалды."""
    monkeypatch.setattr(service.ai, "classify_reply", _ai_resolution(
        solution="Куратор мен оқушы сөйлесті, куратор 5-30 минутта жауап береді", satisfied="Иә"))
    post(env, TECH(), SMART, complaint_text(student="Ламья Жамалова"))
    caption = ("Қайырлы кеш\nШешімі: Негізі куратор 5-30 минут ішінде үнемі уақытылы жауап береді. "
               "Куратор мен оқушы сөйлесті, оқушы шешімге риза")
    m = M(COORD(), SMART, caption=caption, ctype="photo", reply=env.bot.sent[0].id)
    m.photo = [NS(file_id="x")]
    run(bm.handle(m, env.bot, False))
    assert env.bot.reactions[-1][2][0].emoji == "👌"
    q = env.bot.sent[-1]
    assert "негізделген" in q.text and q.kb is not None                  # бот үнсіз қалмады, түймелі сұрақ бар
    assert db.q1("SELECT COUNT(*) n FROM ephemeral")["n"] == 0            # сұрақ өздігінен өшпейді


def test_reply_to_closed_or_deleted_complaint_is_explained_not_ignored(env):
    m = post(env, TECH(), SMART, complaint_text())
    post(env, COORD(), SMART, resolution_text(), reply=m.message_id)            # жабылды
    post(env, COORD(), SMART, resolution_text(), reply=m.message_id)            # қайта жауап
    assert "жабылған" in env.bot.sent[-1].text
    db.update("complaints", C()[0]["id"], status="deleted")
    post(env, COORD(), SMART, "в процессе, шешемін", reply=m.message_id)
    assert "өшірілген" in env.bot.sent[-1].text


def test_status_lists_recent_decisions(env):
    post(env, TECH(), SMART, complaint_text())
    post(env, COORD(), SMART, "в процессе", reply=env.bot.sent[0].id)
    post(env, COORD(), SMART, "ертең жазамын", reply=env.bot.sent[0].id)
    st = service.fmt_status()
    assert "Соңғы жауаптарға бот шешімі" in st and "процессте" in st


# ===================================================================== ағын көрсетілмесе де қабылдау (супервайзерге)
NO_STREAM_CAPTION = """Қайырлы күн!
Жанай Сапар
77718946341

Мәселе: биология кураторын  ауыстырғысы келеді
@mereyka_bio140"""
NO_STREAM_EXT = {"is_complaint": True, "student_name": "Жанай Сапар", "phone": "77718946341", "subject": "БИО",
                 "employee_name": None, "description": "биология кураторын ауыстырғысы келеді"}


def test_screenshot_case_no_stream_is_accepted_and_only_the_supervisor_is_tagged(env, monkeypatch):
    monkeypatch.setattr(service.ai, "extract_complaint", AsyncMock(return_value=NO_STREAM_EXT))
    m = M(TECH(), SMART, caption=NO_STREAM_CAPTION, ctype="photo")
    m.photo = [NS(file_id="x")]
    run(bm.handle(m, env.bot, False))
    c = C()
    assert len(c) == 1 and c[0]["status"] == "open"
    t = env.bot.sent[0].text
    assert "@mereyka_bio140" in t and "Ағын көрсетілмеген" in t and "edit" in t          # супервайзер + нұсқау
    assert "@parvina_bio140" not in t and not any("Түзетіңіз" in x for x in env.bot.texts())
    p = json.loads(c[0]["parsed"])
    assert p["no_stream"] is True and p["subject_label"] == "SMART БИО"


def test_adding_the_stream_by_edit_tags_the_coordinators_without_a_duplicate_complaint(env, monkeypatch):
    monkeypatch.setattr(service.ai, "extract_complaint", AsyncMock(return_value=NO_STREAM_EXT))
    m = M(TECH(), SMART, caption=NO_STREAM_CAPTION, ctype="photo")
    m.photo = [NS(file_id="x")]
    run(bm.handle(m, env.bot, False))
    n = len(env.bot.sent)
    m.caption = NO_STREAM_CAPTION.replace("Мәселе:", "БИО қыркүйек. Мәселе:")                  # менеджер ағынды қосты
    monkeypatch.setattr(service.ai, "extract_complaint", AsyncMock(return_value=dict(NO_STREAM_EXT, subject="БИО қыркүйек")))
    run(bm.handle(m, env.bot, True))
    assert len(C()) == 1 and json.loads(C()[0]["parsed"])["pairs"] == [["БИО", "21"]]
    upd = env.bot.sent[n:]
    assert len(upd) == 1 and "жаңартылды" in upd[0].text and "@ayaulym_bio140" in upd[0].text   # 21 ағын БИО координаторлары


def test_unrecognised_subject_is_still_an_error(env):
    post(env, TECH(), SMART, complaint_text(subject="SMART жоқ пән"))
    assert C() == [] and "танылмады" in env.bot.sent[0].text


# ===================================================================== /update (админ жетіспейтінді қосады)
def _update(env, args, reply_to=None, admin=True):
    out = []

    async def reply(text, **kw):
        out.append(text)
        return NS(chat=NS(id=SMART[0]), message_id=888)
    msg = NS(chat=NS(id=SMART[0]), from_user=NS(id=99 if admin else 5, full_name="Админ"), reply=reply,
             reply_to_message=NS(message_id=reply_to) if reply_to else None)
    run(bm.cmd_update(msg, NS(args=args), env.bot))
    return out


def _no_stream_complaint(env, monkeypatch):
    monkeypatch.setattr(service.ai, "extract_complaint", AsyncMock(return_value=NO_STREAM_EXT))
    m = M(TECH(), SMART, caption=NO_STREAM_CAPTION, ctype="photo")
    m.photo = [NS(file_id="x")]
    run(bm.handle(m, env.bot, False))
    return m


def test_parse_update_args():
    ch, er = service.parse_update_args("ағын=21; пән: БИО\nтелефон 87001112233; Оқушы = Аты Жөні")
    assert ch == {"stream": "21", "subject": "БИО", "phone": "87001112233", "student": "Аты Жөні"} and er == []
    ch, er = service.parse_update_args("бейтаныс=1; ағын 31")
    assert ch == {"stream": "31"} and er == ["«бейтаныс» белгісіз өріс"]


def test_admin_adds_the_missing_stream_and_coordinators_are_tagged(cbot, monkeypatch):
    env = cbot
    _no_stream_complaint(env, monkeypatch)
    cid = C()[0]["id"]
    n = len(env.bot.sent)
    out = _update(env, f"{cid} ағын=қыркүйек")
    assert "жаңартылды" in out[0] and "ағын" in out[0]
    p = json.loads(C()[0]["parsed"])
    assert p["pairs"] == [["БИО", "21"]] and p["no_stream"] is False and p["subject_label"] == "SMART БИО-21"
    upd = env.bot.sent[n:]
    assert len(upd) == 1 and "@ayaulym_bio140" in upd[0].text and "@mereyka_bio140" in upd[0].text
    assert "Ағын көрсетілмеген" not in (C()[0]["notes"] or "")


def test_update_by_reply_and_several_fields_at_once(cbot, monkeypatch):
    env = cbot
    _no_stream_complaint(env, monkeypatch)
    notify = env.bot.sent[0]
    out = _update(env, "ағын=11; қызметкер=Аяқөз Қыдырбай; қызметі=мұғалім; мәселе=куратор ауыстыру", reply_to=notify.id)
    assert "✅" in out[0]
    p = json.loads(C()[0]["parsed"])
    assert (p["pairs"], p["employee_name"], p["employee_role"], p["issue_title"]) == (
        [["БИО", "11"]], "Аяқөз Қыдырбай", "Мұғалім", "Куратор ауыстыру")


def test_update_changes_subject_keeping_the_known_stream(cbot):
    env = cbot
    post(env, TECH(), SMART, complaint_text(subject="SMART ГЕО-31"))
    _update(env, f"{C()[0]['id']} пән=ФИЗ, МАТ")
    p = json.loads(C()[0]["parsed"])
    assert p["pairs"] == [["ФИЗ", "31"], ["МАТ", "31"]] and p["subject_label"] == "SMART ФИЗ-31, МАТ-31"
    assert "@zhetiruzamira" in env.bot.sent[-1].text                         # 31 ағын ФИЗ координаторы жаңадан белгіленді


def test_update_rejects_bad_values_and_unknown_complaints(cbot):
    env = cbot
    post(env, TECH(), SMART, complaint_text())
    cid = C()[0]["id"]
    assert "қызметі" in _update(env, f"{cid} қызметі=ойдан")[0]
    assert "танылмады" in _update(env, f"{cid} пән=жоқ нәрсе")[0]
    assert "ағын" in _update(env, f"{cid} ағын=99")[0]
    assert "Табылмады" in _update(env, "9999 ағын=21")[0]
    assert "Қолданылуы" in _update(env, "")[0]
    assert json.loads(C()[0]["parsed"])["pairs"] == [["ГЕО", "31"]]          # өзгеріссіз қалды


def test_update_requires_admin_and_does_not_retag_when_nothing_new(cbot):
    env = cbot
    post(env, TECH(), SMART, complaint_text())
    cid = C()[0]["id"]
    assert _update(env, f"{cid} ағын=21", admin=False) == []
    assert json.loads(C()[0]["parsed"])["pairs"] == [["ГЕО", "31"]]
    n = len(env.bot.sent)
    _update(env, f"{cid} оқушы=Жаңа Аты; телефон=87001234567")               # жауаптылар өзгермейді
    assert len(env.bot.sent) == n and json.loads(C()[0]["parsed"])["student_name"] == "Жаңа Аты"


def test_update_corrects_a_resolved_complaint_for_the_excel(cbot):
    env = cbot
    m = post(env, TECH(), SMART, complaint_text())
    post(env, COORD(), SMART, resolution_text(), reply=m.message_id)
    n = len(env.bot.sent)
    _update(env, f"{C()[0]['id']} мәселе=Куратор ауыстыру")
    assert len(env.bot.sent) == n                                           # жабылған шағымға хабарлама жоқ
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(exporter.build()))["ТЕХ МЕНЕДЖЕР"]
    assert ws[2][5].value == "Куратор ауыстыру"
