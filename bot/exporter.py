"""Excel экспорт: баған реті түпнұсқа 'ШАҒЫМДАР ЖИНАҒЫ' парақтарымен бірдей."""
import json
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import db
from .config import TZ

TECH_COLS = ["Мәселенің жазылған күні", "Оқушы аты-жөні", "Пәні және потогы", "Қызметкердің аты-жөні",
             "Қызметі/Мәселе иесі", "Мәселе атауы", "Мәселенің толық сипаттамасы", "Білім бөлімінің жауабы",
             "Мәселе шындыққа негізделген бе?", "Оқушымен чат сілтемесі", "Жауапты тех.менеджер",
             "Оқушы жауапқа риза ма?", "Оқушы курста қалды ма?"]
EXT_COLS = ["Мәселенің жазылған күні", "Оқушы аты-жөні", "Пәні және потогы", "Қызметкердің аты-жөні",
            "Қызметі/Мәселе иесі", "Мәселе атауы", "Мәселенің толық сипаттамасы", "Білім бөлімінің жауабы",
            "Мәселеге тренер жауабы", "Мәселе шындыққа негізделген бе?", "Оқушымен чат сілтемесі",
            "Жауапты менеджер", "Оқушы жауапқа риза ма?", "Оқушы курста қалды ма?"]
SHEETS = {"tech": ("ТЕХ МЕНЕДЖЕР", TECH_COLS), "internal": ("ІШКІ САТУ", TECH_COLS),
          "external": ("СЫРТҚЫ САТУ", EXT_COLS), None: ("РӨЛІ БЕЛГІСІЗ", TECH_COLS)}


def _row(c, ext: bool):
    p = json.loads(c["parsed"] or "{}")
    f = json.loads(c["final"] or "{}")
    author = (c["author_name"] or "").split("|")[0].strip()
    dt = datetime.fromisoformat(c["created_at"]).astimezone(TZ).replace(tzinfo=None)
    base = [dt, p.get("student_name") or "-", p.get("subject_label") or "-", p.get("employee_name") or "-",
            p.get("employee_role") or "-", p.get("issue_title") or "-", p.get("description") or c["text"],
            f.get("edu_response") or "-"]
    just = f.get("justified") or ""
    link = p.get("chat_link") or "-"
    sat = f.get("satisfied") or ""
    stay = {"Иә": True, "Жоқ": False}.get(f.get("stayed"), "")
    if ext:
        return base + ["-", just, link, author, sat, stay]
    return base + [just, link, author, sat, stay]


def build(include_open: bool = False, month: str | None = None) -> bytes:
    sql = "SELECT * FROM complaints WHERE status IN ({}) ORDER BY created_at".format(
        "'resolved','open'" if include_open else "'resolved'")
    rows = [c for c in db.q(sql) if not month or c["created_at"].startswith(month)]
    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="D9E1F2")
    for key, (title, cols) in SHEETS.items():
        mine = [c for c in rows if c["author_role"] == key]
        if not mine and key is None:
            continue
        ws = wb.create_sheet(title)
        ws.append(cols)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for c in mine:
            ws.append(_row(c, key == "external"))
        for r in ws.iter_rows(min_row=2):
            r[0].number_format = "DD.MM.YYYY"
            for cell in r:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"  # қолданушы мәтіні формула болып орындалмасын
        widths = [17, 22, 21, 23, 24, 24, 56, 48] + [25] * 6
        for i, w in enumerate(widths[:len(cols)], 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
